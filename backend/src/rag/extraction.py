"""Structure-preserving text extraction for PDF, DOCX, Markdown and plain-text documents.

Each extractor produces an :class:`ExtractedDocument`: an ordered list of :class:`Block`
objects (heading, paragraph, list, table or code) that carry their heading hierarchy,
page references and character offsets into the canonical document text
(``ExtractedDocument.text``, the block texts joined by blank lines). Semantic sections
and chunks refer back to blocks by index, so their text is always rebuilt from the
original extracted content.

Parsing is CPU-bound, so the public entry point :func:`extract_document` runs it in a
worker thread. No OCR is performed: PDFs without an extractable text layer are rejected
with a clear error instead of being stored empty.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import re
import zipfile
from bisect import bisect_right
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any, Literal

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".md", ".markdown", ".txt")
# Below this many characters per page on average, a PDF is treated as scanned/image-only.
MIN_PDF_CHARS_PER_PAGE = 20
BLOCK_SEPARATOR = "\n\n"

BlockKind = Literal["heading", "paragraph", "list", "table", "code"]
TitleSource = Literal["metadata", "heading", "filename"]


class ExtractionError(ValueError):
    """The file could not be turned into text; the message is safe to show to users."""


@dataclass
class Block:
    text: str
    kind: BlockKind = "paragraph"
    level: int = 0  # heading level (0 = document title, 1 = top-level); 0 for content
    heading_path: tuple[str, ...] = ()  # enclosing headings; includes itself for headings
    path_levels: tuple[int, ...] = ()  # heading level of each ``heading_path`` entry
    page: int | None = None
    page_end: int | None = None
    start: int = 0  # character offsets into ExtractedDocument.text
    end: int = 0

    @property
    def heading(self) -> str | None:
        return self.heading_path[-1] if self.heading_path else None

    @property
    def is_heading(self) -> bool:
        return self.kind == "heading"

    @property
    def context_level(self) -> int:
        """Level of the innermost heading this block belongs to (its own, for headings)."""
        return self.path_levels[-1] if self.path_levels else 0

    @property
    def last_page(self) -> int | None:
        return self.page_end if self.page_end is not None else self.page


@dataclass
class ExtractedDocument:
    title: str
    file_type: str
    blocks: list[Block] = field(default_factory=list)
    page_count: int | None = None
    title_source: TitleSource = "filename"
    # Metadata embedded in the file itself (PDF info, DOCX core properties, YAML front
    # matter). Values are JSON-compatible; nothing here is generated.
    source_metadata: dict[str, Any] = field(default_factory=dict)
    text: str = ""

    @property
    def char_count(self) -> int:
        return sum(len(b.text) for b in self.blocks)

    def outline(self) -> list[tuple[int, str, int | None]]:
        """``(level, heading, page)`` for every heading, in document order."""
        return [(b.level, b.text, b.page) for b in self.blocks if b.is_heading]

    def finalize(self) -> ExtractedDocument:
        """Drop empty blocks and assign offsets into the canonical ``text``."""
        self.blocks = [b for b in self.blocks if b.text.strip()]
        offset = 0
        for block in self.blocks:
            block.start = offset
            block.end = offset + len(block.text)
            offset = block.end + len(BLOCK_SEPARATOR)
        self.text = BLOCK_SEPARATOR.join(b.text for b in self.blocks)
        return self


class _StructureBuilder:
    """Tracks the heading stack while an extractor emits blocks in reading order."""

    def __init__(self) -> None:
        self.blocks: list[Block] = []
        self._stack: list[tuple[int, str]] = []

    def heading(self, text: str, level: int, page: int | None = None) -> None:
        text = _clean(text)
        if not text:
            return
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, text))
        self.blocks.append(
            Block(text, "heading", level, *self._path(), page=page)  # type: ignore[arg-type]
        )

    def content(self, text: str, kind: BlockKind = "paragraph", page: int | None = None) -> None:
        text = text.strip("\n").rstrip()
        if not text.strip():
            return
        path, levels = self._path()
        last = self.blocks[-1] if self.blocks else None
        # Consecutive list items under the same heading form one list block.
        if kind == "list" and last and last.kind == "list" and last.heading_path == path:
            last.text = f"{last.text}\n{text}"
            if page is not None and page != last.page:
                last.page_end = page
            return
        self.blocks.append(Block(text, kind, 0, path, levels, page=page))

    def _path(self) -> tuple[tuple[str, ...], tuple[int, ...]]:
        return tuple(t for _, t in self._stack), tuple(lv for lv, _ in self._stack)

    @property
    def first_top_heading(self) -> str | None:
        return next((b.text for b in self.blocks if b.is_heading and b.level <= 1), None)


async def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    return await asyncio.to_thread(extract_document_sync, filename, content)


def extract_document_sync(filename: str, content: bytes) -> ExtractedDocument:
    if not content:
        raise ExtractionError("File is empty")
    file_type = detect_file_type(filename, content)
    stem = PurePath(filename).stem or filename
    if file_type == "pdf":
        doc = _extract_pdf(content, stem)
    elif file_type == "docx":
        doc = _extract_docx(content, stem)
    elif file_type == "md":
        doc = _extract_markdown(_decode_text(content), stem)
    else:
        doc = _extract_text(_decode_text(content), stem)
    doc.finalize()
    if not doc.blocks:
        raise ExtractionError("No extractable text found in document")
    return doc


def _resolve_title(
    metadata_title: Any, builder: _StructureBuilder, stem: str
) -> tuple[str, TitleSource]:
    if isinstance(metadata_title, str) and metadata_title.strip():
        return _clean(metadata_title), "metadata"
    heading = builder.first_top_heading
    if heading:
        return heading, "heading"
    return stem, "filename"


def detect_file_type(filename: str, content: bytes) -> str:
    """Determine the type from the extension and verify it against the file content."""
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'none'}'. Accepted: PDF, DOCX, MD, TXT"
        )
    if ext == ".pdf":
        if not content.lstrip()[:5].startswith(b"%PDF-"):
            raise ExtractionError("File has a .pdf extension but is not a valid PDF")
        return "pdf"
    if ext == ".docx":
        if not _is_docx(content):
            raise ExtractionError("File has a .docx extension but is not a valid Word document")
        return "docx"
    if b"\x00" in content[:8192] and not content.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ExtractionError("File appears to be binary, not text")
    return "md" if ext in (".md", ".markdown") else "txt"


def _is_docx(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            return "word/document.xml" in zf.namelist()
    except zipfile.BadZipFile:
        return False


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _json_value(value: Any) -> Any:
    """Make a metadata value JSON-compatible (dates to ISO strings, recursively)."""
    if isinstance(value, dt.date | dt.datetime):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [_json_value(v) for v in value]
    if value is None or isinstance(value, str | int | float | bool):
        return value
    return str(value)


# Shared heading/list patterns
_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+\S")
_BULLET = re.compile(r"^\s*(?:[•●▪◦‣∙·*+\-–]|\(?\d{1,3}[.)]|\(?[a-z][.)])\s+")
_TERMINAL_PUNCT = (".", "!", "?", ",", ";", ":")


def _numbering_depth(text: str) -> int | None:
    match = _NUMBERED_HEADING.match(text)
    return len(match.group(1).split(".")) if match else None


# --------------------------------------------------------------------------- PDF


@dataclass
class _PdfItem:
    text: str
    bbox: tuple[float, float, float, float]
    size: float = 0.0
    bold: bool = False
    line_count: int = 1
    is_list: bool = False
    is_table: bool = False
    page: int = 1

    @property
    def x0(self) -> float:
        return self.bbox[0]

    @property
    def y0(self) -> float:
        return self.bbox[1]

    @property
    def x1(self) -> float:
        return self.bbox[2]


def _extract_pdf(content: bytes, stem: str) -> ExtractedDocument:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=content, filetype="pdf")
    except Exception as exc:  # pymupdf raises several internal error types
        raise ExtractionError(f"Could not open PDF: {exc}") from exc
    with pdf:
        if pdf.needs_pass:
            raise ExtractionError("PDF is password-protected; upload an unencrypted copy")
        raw_meta = pdf.metadata or {}
        page_count = pdf.page_count
        raw_pages = [
            _RawPage(n, page.rect.width, page.rect.height, page.get_text("dict"), _tables(page))
            for n, page in enumerate(pdf, start=1)
        ]

    skip = _running_line_keys(raw_pages)
    ordered = [
        item for raw in raw_pages for item in _reading_order(_page_items(raw, skip), raw.width)
    ]
    body_size = _dominant_size(ordered)
    tiers = sorted(
        {round(i.size, 1) for i in ordered if not i.is_table and i.size >= body_size * 1.15},
        reverse=True,
    )
    levels = [None if i.is_table else _pdf_heading_level(i, body_size, tiers) for i in ordered]
    # A unique largest heading at the start of page 1 is the document title (level 0).
    headings = [(i, item) for i, item in enumerate(ordered) if levels[i] is not None]
    if len(headings) > 1 and headings[0][1].page == 1:
        first_index, first = headings[0]
        if all(first.size > other.size + 0.5 for _, other in headings[1:]):
            levels[first_index] = 0

    builder = _StructureBuilder()
    for item, level in zip(ordered, levels, strict=True):
        if item.is_table:
            builder.content(item.text, "table", item.page)
        elif level is not None:
            builder.heading(item.text, level, item.page)
        else:
            builder.content(item.text, "list" if item.is_list else "paragraph", item.page)

    chars = sum(len(b.text) for b in builder.blocks)
    if chars < MIN_PDF_CHARS_PER_PAGE * max(page_count, 1):
        raise ExtractionError(
            "PDF has no extractable text layer (it looks scanned or image-only). "
            "OCR is not supported; upload a text-based PDF, DOCX, MD or TXT file."
        )
    source_metadata = {
        key: _clean(str(raw_meta[key]))
        for key in ("title", "author", "subject", "keywords", "creationDate", "modDate")
        if raw_meta.get(key) and str(raw_meta[key]).strip()
    }
    title, title_source = _resolve_title(raw_meta.get("title"), builder, stem)
    return ExtractedDocument(
        title=title,
        file_type="pdf",
        blocks=builder.blocks,
        page_count=page_count,
        title_source=title_source,
        source_metadata=source_metadata,
    )


@dataclass
class _RawPage:
    number: int
    width: float
    height: float
    text: dict
    tables: list[_PdfItem]


def _tables(page: Any) -> list[_PdfItem]:
    """Tables found by PyMuPDF's ruling-line detection, rendered as ``a | b`` rows."""
    items: list[_PdfItem] = []
    try:
        for table in page.find_tables().tables:
            rows = [
                " | ".join(_clean(cell) for cell in row if cell and _clean(cell))
                for row in table.extract()
            ]
            text = "\n".join(r for r in rows if r)
            if text:
                items.append(_PdfItem(text, tuple(table.bbox), is_table=True))
    except Exception:  # table detection is best-effort; text blocks are still extracted
        return []
    return items


def _in_margin(bbox: tuple[float, ...], height: float) -> bool:
    return bbox[3] <= height * 0.08 or bbox[1] >= height * 0.92


def _line_key(text: str) -> str:
    return re.sub(r"\d+", "#", " ".join(text.split()).casefold())


def _line_text(line: dict) -> str:
    return _clean("".join(s.get("text", "") for s in line.get("spans", [])))


_PAGE_NUMBER = re.compile(r"^(page\s*)?\d{1,4}(\s*(of|/)\s*\d{1,4})?$", re.IGNORECASE)


def _running_line_keys(pages: list[_RawPage]) -> set[str]:
    """Header/footer lines repeated in the page margins of most pages."""
    if len(pages) < 3:
        return set()
    counts: dict[str, int] = {}
    for page in pages:
        keys = {
            _line_key(text)
            for block in page.text.get("blocks", [])
            for line in block.get("lines", [])
            if (text := _line_text(line)) and _in_margin(line["bbox"], page.height)
        }
        for key in keys:
            counts[key] = counts.get(key, 0) + 1
    threshold = max(3, len(pages) // 2)
    return {key for key, count in counts.items() if count >= threshold}


def _page_items(page: _RawPage, skip: set[str]) -> list[_PdfItem]:
    """Text blocks and tables of one page, without running headers/footers."""
    items = [_PdfItem(t.text, t.bbox, is_table=True, page=page.number) for t in page.tables]
    boxes = [t.bbox for t in page.tables]
    for block in page.text.get("blocks", []):
        if block.get("type") != 0:  # 0 = text, 1 = image
            continue
        lines, sizes, bold = [], [], True
        box: tuple[float, float, float, float] | None = None
        for line in block.get("lines", []):
            spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
            text = _line_text(line)
            if not spans or not text:
                continue
            x0, y0, x1, y1 = line["bbox"]
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            if any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes):
                continue  # already part of a detected table
            if _in_margin(line["bbox"], page.height) and (
                _line_key(text) in skip or _PAGE_NUMBER.match(text)
            ):
                continue
            lines.append(text)
            sizes.extend(s["size"] for s in spans)
            bold = bold and all(s.get("flags", 0) & 16 for s in spans)
            if box is None:
                box = (x0, y0, x1, y1)
            else:
                box = (min(box[0], x0), min(box[1], y0), max(box[2], x1), max(box[3], y1))
        if not lines or box is None:
            continue
        bullets = sum(bool(_BULLET.match(line)) for line in lines)
        is_list = bullets > 0 and _BULLET.match(lines[0]) is not None
        text = "\n".join(lines) if is_list and bullets > 1 else _join_lines(lines)
        items.append(_PdfItem(text, box, max(sizes), bold, len(lines), is_list, page=page.number))
    return items


def _join_lines(lines: list[str]) -> str:
    text = ""
    for line in lines:
        if text.endswith("-") and len(text) > 1 and text[-2].isalpha() and line[:1].islower():
            text = text[:-1] + line  # re-join a word hyphenated across a line break
        else:
            text = f"{text} {line}" if text else line
    return text


def _reading_order(items: list[_PdfItem], width: float) -> list[_PdfItem]:
    """Order blocks top-to-bottom, reading two-column regions column by column.

    PyMuPDF's own ordering follows the content stream (or pure y/x position with
    ``sort=True``), which interleaves columns. Blocks that span the middle of the page
    (titles, full-width paragraphs, tables) split the page into bands; inside each band
    the left column is read before the right one.
    """
    if not items:
        return []
    mid, gutter = width / 2, width * 0.04
    left = [i for i in items if i.x1 <= mid + gutter and i.x0 < mid - gutter]
    right = [i for i in items if i.x0 >= mid - gutter and i.x1 > mid + gutter]
    total = sum(len(i.text) for i in items) or 1
    two_columns = (
        len(left) >= 2
        and len(right) >= 2
        and sum(len(i.text) for i in left) / total > 0.15
        and sum(len(i.text) for i in right) / total > 0.15
    )
    if not two_columns:
        return sorted(items, key=lambda i: (round(i.y0), i.x0))
    column_ids = {id(i) for i in left} | {id(i) for i in right}
    spanning = sorted((i for i in items if id(i) not in column_ids), key=lambda i: i.y0)
    starts = [s.y0 for s in spanning]
    bands: list[tuple[list[_PdfItem], list[_PdfItem]]] = [
        ([], []) for _ in range(len(spanning) + 1)
    ]
    for column, target in ((left, 0), (right, 1)):
        for item in column:
            bands[bisect_right(starts, item.y0)][target].append(item)
    ordered: list[_PdfItem] = []
    for index, (band_left, band_right) in enumerate(bands):
        if index > 0:
            ordered.append(spanning[index - 1])
        ordered.extend(sorted(band_left, key=lambda i: i.y0))
        ordered.extend(sorted(band_right, key=lambda i: i.y0))
    return ordered


def _dominant_size(items: list[_PdfItem]) -> float:
    weights: dict[float, int] = {}
    for item in items:
        if not item.is_table:
            size = round(item.size, 1)
            weights[size] = weights.get(size, 0) + len(item.text)
    return max(weights, key=lambda s: weights[s]) if weights else 11.0


def _pdf_heading_level(item: _PdfItem, body_size: float, tiers: list[float]) -> int | None:
    text = item.text
    if len(text) > 120 or item.line_count > 2 or text.endswith(_TERMINAL_PUNCT):
        return None
    depth = _numbering_depth(text)
    size = round(item.size, 1)
    if size >= body_size * 1.15:  # larger font: a heading even if it looks numbered
        return depth or (tiers.index(size) + 1 if size in tiers else 1)
    if item.is_list and not item.bold:
        return None
    if item.bold and size >= body_size and len(text) <= 80:
        return depth or len(tiers) + 1
    if depth is not None and depth >= 2 and len(text) <= 80:  # "2.1 VPN Configuration"
        return depth
    return None


# -------------------------------------------------------------------------- DOCX


def _extract_docx(content: bytes, stem: str) -> ExtractedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(f"Could not open Word document: {exc}") from exc

    builder = _StructureBuilder()
    numbered_run = 0
    # iter_inner_content yields paragraphs and tables in their original document order.
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = _clean(item.text)
            if not text:
                continue
            style = (item.style.name if item.style is not None else "") or ""
            level = _docx_heading_level(style)
            if level is not None:
                builder.heading(text, level)
                numbered_run = 0
                continue
            p_pr = item._p.pPr
            is_list = style.startswith("List") or (p_pr is not None and p_pr.numPr is not None)
            if is_list:
                numbered_run += 1
                marker = f"{numbered_run}." if "Number" in style else "-"
                builder.content(f"{marker} {text}", "list")
            else:
                numbered_run = 0
                builder.content(text)
        elif isinstance(item, Table):
            numbered_run = 0
            rows = []
            for row in item.rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = _clean(cell.text)
                    if value and (not cells or cells[-1] != value):  # skip merged duplicates
                        cells.append(value)
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                builder.content("\n".join(rows), "table")

    props = document.core_properties
    source_metadata = {
        name: _json_value(value)
        for name in (
            "title",
            "author",
            "subject",
            "keywords",
            "category",
            "comments",
            "version",
            "identifier",
            "created",
            "modified",
        )
        if (value := getattr(props, name, None)) not in (None, "")
    }
    title, title_source = _resolve_title(props.title, builder, stem)
    return ExtractedDocument(
        title=title,
        file_type="docx",
        blocks=builder.blocks,
        title_source=title_source,
        source_metadata=source_metadata,
    )


def _docx_heading_level(style: str) -> int | None:
    if style == "Title":
        return 0
    match = re.fullmatch(r"Heading\s*(\d)", style)
    return int(match.group(1)) if match else None


# ---------------------------------------------------------------------- Markdown

_ATX = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_SETEXT = re.compile(r"^\s{0,3}(=+|-+)\s*$")
_FENCE = re.compile(r"^\s{0,3}(```|~~~)")
_MD_LIST = re.compile(r"^\s{0,3}(?:[-*+]|\d{1,3}[.)])\s+")
_MD_TABLE = re.compile(r"^\s*\|")


def _extract_markdown(text: str, stem: str) -> ExtractedDocument:
    import frontmatter

    try:
        post = frontmatter.loads(text)
        body, meta = post.content, _json_value(dict(post.metadata))
    except Exception:  # malformed front matter: treat the whole file as body
        body, meta = text, {}

    builder = _StructureBuilder()
    paragraph: list[str] = []
    in_fence = False

    def flush() -> None:
        if not paragraph:
            return
        joined = "\n".join(paragraph).strip("\n")
        first = paragraph[0]
        if _FENCE.match(first):
            kind: BlockKind = "code"
        elif _MD_TABLE.match(first):
            kind = "table"
        elif _MD_LIST.match(first):
            kind = "list"
        else:
            kind = "paragraph"
        builder.content(joined, kind)
        paragraph.clear()

    for line in body.splitlines():
        if _FENCE.match(line):
            if not in_fence:
                flush()
            in_fence = not in_fence
            paragraph.append(line)
            if not in_fence:
                flush()
            continue
        if in_fence:
            paragraph.append(line)
            continue
        atx = _ATX.match(line)
        if atx:
            flush()
            builder.heading(atx.group(2), len(atx.group(1)))
            continue
        setext = _SETEXT.match(line)
        if setext and len(paragraph) == 1 and paragraph[0].strip():
            heading = paragraph.pop()
            builder.heading(heading, 1 if setext.group(1).startswith("=") else 2)
            continue
        if not line.strip():
            flush()
            continue
        if paragraph and _MD_LIST.match(line) and not _MD_LIST.match(paragraph[0]):
            flush()  # a list directly after a paragraph line starts a new block
        paragraph.append(line.rstrip())
    flush()
    title, title_source = _resolve_title(meta.get("title"), builder, stem)
    return ExtractedDocument(
        title=title,
        file_type="md",
        blocks=builder.blocks,
        title_source=title_source,
        source_metadata=meta,
    )


# -------------------------------------------------------------------------- Text

_UNDERLINE = re.compile(r"^\s*(=+|-+)\s*$")


def _extract_text(text: str, stem: str) -> ExtractedDocument:
    """Plain text has no markup, so only unambiguous heading patterns are recognised:
    numbered headings ("2.1 VPN Setup"), ALL-CAPS lines and underlined lines."""
    builder = _StructureBuilder()
    for raw in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [line.rstrip() for line in raw.strip("\n").splitlines() if line.strip()]
        if not lines:
            continue
        if len(lines) == 2 and _UNDERLINE.match(lines[1]) and len(lines[0].strip()) <= 100:
            builder.heading(lines[0], 1 if lines[1].strip().startswith("=") else 2)
            continue
        if len(lines) == 1 and (level := _text_heading_level(lines[0].strip())):
            builder.heading(lines[0], level)
            continue
        kind: BlockKind = "list" if _BULLET.match(lines[0]) else "paragraph"
        builder.content("\n".join(lines), kind)
    title, title_source = _resolve_title(None, builder, stem)
    return ExtractedDocument(
        title=title, file_type="txt", blocks=builder.blocks, title_source=title_source
    )


def _text_heading_level(line: str) -> int | None:
    if len(line) > 100 or line.endswith(_TERMINAL_PUNCT):
        return None
    depth = _numbering_depth(line)
    if depth is not None and (depth >= 2 or line.split(maxsplit=1)[0].endswith(".")):
        words = line.split()
        if len(words) <= 10 and words[1][:1].isupper():
            return depth
    letters = [c for c in line if c.isalpha()]
    if len(letters) >= 3 and all(c.isupper() for c in letters) and len(line.split()) <= 10:
        return 1
    return None


def _clean(text: str) -> str:
    return re.sub(r"[ \t ]+", " ", text).strip()
