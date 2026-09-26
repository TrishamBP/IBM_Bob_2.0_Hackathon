"""Structure-preserving extraction (PDF/DOCX/MD/TXT)."""

import io

import docx
import pymupdf
import pytest

from src.rag.extraction import ExtractionError, extract_document

HEADER = "ACME Corp - Internal"


def make_pdf(pages: list[list[tuple[str, float, bool]]], header: bool = False) -> bytes:
    """Each line is (text, font size, bold); lines are stacked from the top margin."""
    pdf = pymupdf.open()
    for number, lines in enumerate(pages, start=1):
        page = pdf.new_page()
        if header:
            page.insert_text((72, 30), HEADER, fontsize=9)
            page.insert_text((290, 820), f"Page {number} of {len(pages)}", fontsize=9)
        y = 90
        for text, size, bold in lines:
            page.insert_text((72, y), text, fontsize=size, fontname="hebo" if bold else "helv")
            y += size * 2.2
    data = pdf.tobytes()
    pdf.close()
    return data


BODY = 11
GUIDE_PAGES = [
    [
        ("IT Employee Onboarding Guide", 22, True),
        ("1. Development Environment", 16, True),
        ("1.1 VS Code Installation", 13, True),
        ("Download VS Code from the Software Center.", BODY, False),
    ],
    [
        ("2. Corporate Network", 16, True),
        ("2.1 VPN Configuration", 13, True),
        ("Prerequisites", BODY, True),
        ("You need an active ACME account and MFA enrolled.", BODY, False),
    ],
    [
        ("2.2 Network Troubleshooting", 13, True),
        ("Restart the client if the VPN drops.", BODY, False),
    ],
]


async def test_pdf_heading_hierarchy_pages_and_running_headers():
    doc = await extract_document("guide.pdf", make_pdf(GUIDE_PAGES, header=True))
    assert doc.file_type == "pdf" and doc.page_count == 3
    assert doc.title == "IT Employee Onboarding Guide"
    assert HEADER not in doc.text and "Page 2 of 3" not in doc.text
    levels = {text: level for level, text, _ in doc.outline()}
    assert levels["IT Employee Onboarding Guide"] == 0
    assert levels["1. Development Environment"] == 1  # numbered, not a list item
    assert levels["1.1 VS Code Installation"] == 2
    assert levels["Prerequisites"] > levels["2.1 VPN Configuration"]
    mfa = next(b for b in doc.blocks if "MFA" in b.text)
    assert mfa.page == 2 and mfa.kind == "paragraph"
    assert mfa.heading_path[-3:] == (
        "2. Corporate Network",
        "2.1 VPN Configuration",
        "Prerequisites",
    )
    restart = next(b for b in doc.blocks if "Restart" in b.text)
    assert restart.page == 3 and restart.heading == "2.2 Network Troubleshooting"


async def test_pdf_two_column_reading_order():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 60), "Benefits Overview", fontsize=20, fontname="hebo")
    left = ["Left column starts here.", "Left column continues.", "Left column ends."]
    right = ["Right column starts here.", "Right column continues.", "Right column ends."]
    filler = " The benefits portal explains every plan option in detail."
    for i, (a, b) in enumerate(zip(left, right, strict=True)):
        top = 110 + i * 200
        page.insert_textbox(pymupdf.Rect(50, top, 280, top + 150), a + filler * 3, fontsize=BODY)
        page.insert_textbox(pymupdf.Rect(320, top, 550, top + 150), b + filler * 3, fontsize=BODY)
    content = pdf.tobytes()
    pdf.close()
    doc = await extract_document("two-col.pdf", content)
    texts = [b.text for b in doc.blocks]
    order = [next(i for i, t in enumerate(texts) if s in t) for s in left + right]
    assert order == sorted(order), texts  # the whole left column precedes the right one


async def test_pdf_offsets_match_canonical_text():
    doc = await extract_document("guide.pdf", make_pdf(GUIDE_PAGES))
    for block in doc.blocks:
        assert doc.text[block.start : block.end] == block.text


async def test_scanned_pdf_requires_ocr_error():
    pdf = pymupdf.open()
    pdf.new_page()
    content = pdf.tobytes()
    with pytest.raises(ExtractionError, match="OCR is not supported"):
        await extract_document("scan.pdf", content)


def make_docx() -> bytes:
    document = docx.Document()
    document.core_properties.title = "Laptop Policy"
    document.core_properties.author = "IT Asset Team"
    document.add_heading("Eligibility", level=1)
    document.add_paragraph("All full-time employees receive a laptop on day one.")
    document.add_heading("Approved Models", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Role", "Model"
    table.cell(1, 0).text, table.cell(1, 1).text = "Engineer", "MacBook Pro"
    document.add_heading("Returns", level=1)
    document.add_paragraph("Wipe personal files.", style="List Bullet")
    document.add_paragraph("Return the charger.", style="List Bullet")
    document.add_paragraph("Return devices to IT within five days of leaving.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_docx_headings_tables_and_lists_in_order():
    doc = await extract_document("policy.docx", make_docx())
    assert doc.title == "Laptop Policy" and doc.title_source == "metadata"
    assert doc.source_metadata["author"] == "IT Asset Team"
    assert [(level, text) for level, text, _ in doc.outline()] == [
        (1, "Eligibility"),
        (2, "Approved Models"),
        (1, "Returns"),
    ]
    table = next(b for b in doc.blocks if b.kind == "table")
    assert "Role | Model" in table.text and "Engineer | MacBook Pro" in table.text
    assert table.heading_path == ("Eligibility", "Approved Models")
    lists = [b for b in doc.blocks if b.kind == "list"]
    assert len(lists) == 1 and "Wipe personal files." in lists[0].text
    assert "Return the charger." in lists[0].text and lists[0].heading == "Returns"
    texts = [b.text for b in doc.blocks]
    assert texts.index(table.text) < texts.index(lists[0].text)


async def test_markdown_hierarchy_front_matter_lists_and_code():
    md = (
        "---\ntitle: Onboarding Guide\nowner: People Team\nversion: 2.1\n---\n"
        "# Week One\n\nMeet your buddy.\n\n"
        "## Accounts\n\n- Laptop login\n- Email\n\n"
        "### Slack\n\nJoin #general.\n\n"
        "```\n# not a heading\n```\n\n"
        "Setext Heading\n--------------\n\nFill in the tax forms.\n\n"
        "| Form | Due |\n|---|---|\n| W-4 | Day 1 |\n"
    )
    doc = await extract_document("guide.md", md.encode())
    assert doc.title == "Onboarding Guide"
    assert doc.source_metadata["owner"] == "People Team"
    by_text = {b.text: b for b in doc.blocks}
    assert by_text["Join #general."].heading_path == ("Week One", "Accounts", "Slack")
    assert by_text["- Laptop login\n- Email"].kind == "list"
    code = by_text["```\n# not a heading\n```"]
    assert code.kind == "code" and code.heading == "Slack"
    tax = by_text["Fill in the tax forms."]
    assert tax.heading_path == ("Week One", "Setext Heading")  # setext "---" is level 2
    assert next(b for b in doc.blocks if b.kind == "table").heading == "Setext Heading"


async def test_txt_headings():
    text = (
        "EXPENSE POLICY\n\nThis policy covers travel.\n\n"
        "1. Submitting Claims\n\nUse the finance portal.\n\n"
        "1.1 Receipts\n\nAttach every receipt.\n"
    )
    doc = await extract_document("policy.txt", text.encode())
    outline = [(level, t) for level, t, _ in doc.outline()]
    assert (1, "1. Submitting Claims") in outline and (2, "1.1 Receipts") in outline
    receipts = next(b for b in doc.blocks if b.text == "Attach every receipt.")
    assert receipts.heading_path[-2:] == ("1. Submitting Claims", "1.1 Receipts")


@pytest.mark.parametrize(
    "content",
    ["Café policy\n\nSecond paragraph".encode(), "Café".encode("cp1252"), b"\xef\xbb\xbfBOM text"],
)
async def test_txt_decoding(content):
    doc = await extract_document("notes.txt", content)
    assert doc.blocks and doc.file_type == "txt"
    assert not doc.blocks[0].text.startswith("﻿")


@pytest.mark.parametrize(
    "filename, content, match",
    [
        ("image.png", b"\x89PNG", "Unsupported file type"),
        ("fake.pdf", b"hello", "not a valid PDF"),
        ("fake.docx", b"PK\x03\x04garbage", "not a valid Word"),
        ("binary.txt", b"abc\x00def", "binary"),
        ("empty.txt", b"", "empty"),
        ("blank.md", b"   \n\n  ", "No extractable text"),
    ],
)
async def test_invalid_files(filename, content, match):
    with pytest.raises(ExtractionError, match=match):
        await extract_document(filename, content)
